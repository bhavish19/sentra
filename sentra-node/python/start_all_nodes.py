"""SENTRA launcher (MNIST-only, batched secure runner).

This launcher is the canonical multi-node entrypoint:
  start_all_nodes.py -> run_mnist_batched_secure.py (per node)

CLI parsing/validation lives in `start_all_nodes_cli.py`.
"""

from __future__ import annotations

import datetime
import json
import re
import subprocess
import sys
import time
<<<<<<< HEAD
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

=======
import os
import json
import datetime
import re
from pathlib import Path
import math
>>>>>>> origin/main
from start_all_nodes_cli import parse_and_validate_args

try:
    import psutil
except Exception:  # pragma: no cover
    psutil = None

ProcEntry = Tuple[int, subprocess.Popen]


<<<<<<< HEAD
def _build_node_command(args, node_id: int) -> List[str]:
    cmd: List[str] = [
        sys.executable,
        "-u",
        "run_mnist_batched_secure.py",
        "--node-id",
        str(int(node_id)),
        "--n-nodes",
        str(int(args.n_nodes)),
        "--t",
        str(int(args.t)),
        "--enable-network",
        "--base-port",
        str(int(args.base_port)),
        "--host",
        str(args.host),
        "--batch-size",
        str(int(args.batch_size)),
        "--accum-steps",
        str(int(args.accum_steps)),
        "--num-epochs",
        str(int(args.num_epochs)),
        "--learning-rate",
        str(float(args.learning_rate)),
        "--mnist-samples",
        str(int(args.mnist_samples)),
        "--seed",
        str(int(args.seed)),
        "--loss-mode",
        str(args.loss_mode),
        "--scale-factor",
        str(int(args.scale_factor)),
        "--field-size",
        str(int(args.field_size)),
        "--softmax-temperature",
        str(float(args.softmax_temperature)),
        "--grad-clip",
        str(float(args.grad_clip)),
        "--logit-clip",
        str(float(args.logit_clip)),
        "--exp-approx",
        str(args.exp_approx),
        "--softmax-grad-mode",
        str(args.softmax_grad_mode),
        "--explode-logit-threshold",
        str(float(args.explode_logit_threshold)),
        "--loss-growth-threshold",
        str(float(args.loss_growth_threshold)),
        "--grad-norm-threshold",
        str(float(args.grad_norm_threshold)),
        "--membership-epoch",
        str(int(args.membership_epoch)),
    ]

    if bool(args.no_abort_on_instability):
        cmd.append("--no-abort-on-instability")
    if bool(args.debug_numerics):
        cmd.append("--debug-numerics")
    if bool(args.debug_division):
        cmd.append("--debug-division")
    if bool(args.packed_forward_pilot):
        cmd.append("--packed-forward-pilot")
    if bool(args.packed_forward_native):
        cmd.append("--packed-forward-native")
    if bool(args.packed_end2end):
        cmd.append("--packed-end2end")
    if int(args.dpss_refresh_interval) > 0:
        cmd.extend(["--dpss-refresh-interval", str(int(args.dpss_refresh_interval))])

    if bool(args.enable_failure_detection):
        cmd.append("--enable-failure-detection")
    if bool(getattr(args, "enable_dropout_reshare_recovery", False)):
        cmd.append("--enable-dropout-reshare-recovery")
    if bool(getattr(args, "enable_join_recovery", False)):
        cmd.append("--enable-join-recovery")

    if bool(args.distribute_dataset_shares):
        cmd.extend(
            [
                "--distribute-dataset-shares",
                "--dataset-owner-node",
                str(int(args.dataset_owner_node)),
                "--dataset-distribution-timeout",
                str(float(args.dataset_distribution_timeout)),
            ]
        )
    if bool(args.receive_dataset_shares_from_client):
        cmd.extend(
            [
                "--receive-dataset-shares-from-client",
                "--dataset-source-node-id",
                str(int(args.dataset_source_node_id)),
                "--dataset-distribution-timeout",
                str(float(args.dataset_distribution_timeout)),
            ]
        )
        if bool(args.client_eval_after_training):
            cmd.extend(
                [
                    "--client-eval-after-training",
                    "--client-eval-samples",
                    str(int(args.client_eval_samples)),
                ]
            )

    if bool(args.export_reconstructed_model):
        cmd.extend(
            [
                "--export-reconstructed-model",
                str(args.export_reconstructed_model),
                "--export-timeout",
                str(float(args.export_timeout)),
            ]
        )

    if bool(getattr(args, "use_kvs_dataset", False)):
        cmd.append("--use-kvs-dataset")
    if bool(getattr(args, "use_weight_versioning", False)):
        cmd.append("--use-weight-versioning")

=======
def _build_node_command(args, node_id):
    """Build command for a node process."""
    cmd = [
        sys.executable,
        '-u',
        'run_mnist_batched_secure.py',
        '--node-id', str(node_id),
        '--n-nodes', str(args.n_nodes),
        '--base-port', str(args.base_port),
        '--host', str(args.host),
        '--batch-size', str(args.batch_size),
        '--num-epochs', str(args.num_epochs),
        '--learning-rate', str(args.learning_rate),
        '--t', str(args.t),
        '--seed', str(args.seed),
        '--mnist-samples', str(args.mnist_samples),
        '--enable-network',
    ]
    cmd.extend([
        '--accum-steps', str(args.accum_steps),
        '--loss-mode', str(args.loss_mode),
        '--scale-factor', str(args.scale_factor),
        '--field-size', str(args.field_size),
        '--softmax-temperature', str(args.softmax_temperature),
        '--grad-clip', str(args.grad_clip),
        '--logit-clip', str(args.logit_clip),
        '--exp-approx', str(args.exp_approx),
        '--softmax-grad-mode', str(args.softmax_grad_mode),
        '--explode-logit-threshold', str(args.explode_logit_threshold),
        '--loss-growth-threshold', str(args.loss_growth_threshold),
        '--grad-norm-threshold', str(args.grad_norm_threshold),
    ])
    if args.distribute_dataset_shares:
        cmd.extend([
            '--distribute-dataset-shares',
            '--dataset-owner-node', str(args.dataset_owner_node),
            '--dataset-distribution-timeout', str(args.dataset_distribution_timeout),
        ])
    if args.receive_dataset_shares_from_client:
        cmd.extend([
            '--receive-dataset-shares-from-client',
            '--dataset-source-node-id', str(args.dataset_source_node_id),
            '--dataset-distribution-timeout', str(args.dataset_distribution_timeout),
        ])
        if args.client_eval_after_training:
            cmd.extend([
                '--client-eval-after-training',
                '--client-eval-samples', str(args.client_eval_samples),
            ])
    if args.export_reconstructed_model:
        cmd.extend([
            '--export-reconstructed-model', str(args.export_reconstructed_model),
            '--export-timeout', str(args.export_timeout),
        ])
    if args.no_abort_on_instability:
        cmd.append('--no-abort-on-instability')
    if args.debug_numerics:
        cmd.append('--debug-numerics')
    if args.debug_division:
        cmd.append('--debug-division')
    if args.packed_forward_pilot:
        cmd.append('--packed-forward-pilot')
    if args.packed_forward_native:
        cmd.append('--packed-forward-native')
    if args.packed_end2end:
        cmd.append('--packed-end2end')
    if args.dpss_refresh_interval > 0:
        cmd.extend(['--dpss-refresh-interval', str(args.dpss_refresh_interval)])
    cmd.extend(['--membership-epoch', str(args.membership_epoch)])
    if args.enable_failure_detection:
        cmd.append('--enable-failure-detection')
    if getattr(args, "enable_dropout_reshare_recovery", False):
        cmd.append('--enable-dropout-reshare-recovery')
    if getattr(args, "enable_join_recovery", False):
        cmd.append('--enable-join-recovery')
    if getattr(args, "use_kvs_dataset", False):
        cmd.append('--use-kvs-dataset')
    if getattr(args, "use_weight_versioning", False):
        cmd.append('--use-weight-versioning')
>>>>>>> origin/main
    return cmd


def _extract_metrics_from_log(log_text: str) -> Dict[str, Any]:
    metrics: Dict[str, Any] = {}
    # Put '-' at the end of the character class to avoid range parsing issues.
    acc = re.findall(r"Epoch\\s+(\\d+)\\s+Test Accuracy:\\s+([0-9eE+.\\-]+)%", log_text)
    loss = re.findall(r"Epoch\\s+(\\d+)\\s+Test Loss:\\s+([0-9eE+.\\-]+)", log_text)
    if acc:
        _ep, _val = acc[-1]
        metrics["final_epoch_acc_pct"] = float(_val)
    if loss:
        _ep, _val = loss[-1]
        metrics["final_epoch_loss"] = float(_val)
    return metrics


def _append_run_to_xlsx(path: str, row: Dict[str, Any]) -> str:
    try:
        from openpyxl import Workbook, load_workbook  # type: ignore
    except Exception:
        return path
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    if p.exists():
        wb = load_workbook(p)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.append(list(row.keys()))
    ws.append([row.get(k, "") for k in row.keys()])
    wb.save(p)
    return str(p)


def _run_headless(args) -> None:
    run_id = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    run_dir = Path("logs") / f"run_{run_id}"
    run_dir.mkdir(parents=True, exist_ok=True)

    procs: List[ProcEntry] = []
    open_files = []
    for nid in range(1, int(args.n_nodes) + 1):
        out_path = run_dir / f"node_{nid}.log"
        err_path = run_dir / f"node_{nid}.err"
        out_f = open(out_path, "w", encoding="utf-8")
        err_f = open(err_path, "w", encoding="utf-8")
        open_files.extend([out_f, err_f])
        p = subprocess.Popen(_build_node_command(args, nid), cwd=str(Path(".")), stdout=out_f, stderr=err_f)
        procs.append((nid, p))
        print(f"[headless] started node {nid} -> {out_path}")
        time.sleep(0.4)

    client_proc: Optional[subprocess.Popen] = None
    client_log = run_dir / "client_distributor.log"
    if bool(getattr(args, "start_client_distributor", False)):
        cf = open(client_log, "w", encoding="utf-8")
        open_files.append(cf)
        cmd = [
            sys.executable,
            "-u",
            "client_distributor.py",
            "--client-node-id",
            str(int(getattr(args, "dataset_source_node_id", 0))),
            "--n-nodes",
            str(int(args.n_nodes)),
            "--t",
            str(int(args.t)),
            "--base-port",
            str(int(args.base_port)),
            "--host",
            str(args.host),
            "--mnist-samples",
            str(int(args.mnist_samples)),
            "--client-test-samples",
            str(int(args.client_test_samples)),
            "--field-size",
            str(int(args.field_size)),
            "--scale-factor",
            str(int(args.scale_factor)),
            "--seed",
            str(int(args.seed)),
            "--barrier-timeout",
            str(float(args.dataset_distribution_timeout)),
            "--membership-epoch",
            str(int(args.membership_epoch)),
        ]
        if bool(args.client_eval_after_training):
            cmd.extend(
                [
                    "--collect-client-eval",
                    "--client-eval-samples",
                    str(int(args.client_eval_samples)),
                    "--eval-timeout",
                    str(float(args.client_eval_timeout)),
                ]
            )
        client_proc = subprocess.Popen(cmd, cwd=str(Path(".")), stdout=cf, stderr=cf)
        print(f"[headless] started client distributor -> {client_log}")

    t0 = time.time()
    return_codes: Dict[int, int] = {}
    for nid, p in procs:
        p.wait()
        return_codes[nid] = int(p.returncode or 0)
        print(f"[headless] node {nid} exited with code {return_codes[nid]}")
    if client_proc is not None:
        try:
            client_proc.wait(timeout=15.0)
        except Exception:
            pass

    for f in open_files:
        try:
            f.close()
        except Exception:
            pass

    duration = float(time.time() - t0)
    status = "success" if all(c == 0 for c in return_codes.values()) else "failed"

    node1_text = (run_dir / "node_1.log").read_text(encoding="utf-8", errors="ignore")
    m1 = _extract_metrics_from_log(node1_text)
    client_text = client_log.read_text(encoding="utf-8", errors="ignore") if client_log.exists() else ""
    client_acc = None
    # Client prints: "Client Final Accuracy (100 samples): 16.00%"
    # Keep '-' escaped (or last) to avoid regex range parsing.
    m_client = re.findall(r"Client Final Accuracy \(\d+ samples\):\s*([0-9eE+.\\-]+)%", client_text)
    if m_client:
        client_acc = float(m_client[-1])

    print("\nRun summary:")
    print(f"  status: {status}")
    print(f"  final_epoch_acc_pct: {m1.get('final_epoch_acc_pct', 0.0)}")
    print(f"  final_epoch_loss: {m1.get('final_epoch_loss', 0.0)}")
    print(f"  reconstructed_acc: {client_acc if client_acc is not None else ''}")
    print(f"  duration_sec: {duration:.2f}")
    print(f"  logs: {run_dir}")

    if bool(getattr(args, "record_results_xlsx", "")):
        row = {
            "timestamp": datetime.datetime.now().isoformat(timespec="seconds"),
            "status": status,
            "duration_sec": duration,
<<<<<<< HEAD
            "n_nodes": int(args.n_nodes),
            "dataset": str(args.dataset),
            "train_mode": "secure",
            "loss_mode": str(args.loss_mode),
            "softmax_grad_mode": str(args.softmax_grad_mode),
            "num_epochs": int(args.num_epochs),
            "batch_size": int(args.batch_size),
            "learning_rate": float(args.learning_rate),
            "field_size": int(args.field_size),
            "scale_factor": int(args.scale_factor),
            "softmax_temperature": float(args.softmax_temperature),
            "seed": int(args.seed),
            "final_epoch_acc_pct": m1.get("final_epoch_acc_pct", ""),
            "final_epoch_loss": m1.get("final_epoch_loss", ""),
            "reconstructed_acc": client_acc if client_acc is not None else "",
=======
            "duration_min": round(duration / 60.0, 3),
            "n_nodes": args.n_nodes,
            "dataset": args.dataset,
            "batched": bool(args.batched),
            "train_mode": "secure",
            "loss_mode": args.loss_mode,
            "softmax_grad_mode": args.softmax_grad_mode,
            "num_epochs": args.num_epochs,
            "batch_size": args.batch_size,
            "learning_rate": args.learning_rate,
            "field_size": args.field_size,
            "scale_factor": args.scale_factor,
            "softmax_temperature": args.softmax_temperature,
            "grad_clip": args.grad_clip,
            "logit_clip": args.logit_clip,
            "seed": args.seed,
            "final_epoch_acc_pct": chosen.get("last_epoch_acc_pct", ""),
            "final_epoch_loss": chosen.get("last_epoch_loss", ""),
            "reconstructed_acc": reconstructed_acc if reconstructed_acc is not None else "",
            "node_success_count": success_count,
            "node_elapsed_sec": json.dumps(node_elapsed),
            "prover_time_sec": prover_time_sec if prover_time_sec is not None else "",
            "avg_batch_time_sec": avg_batch_time_sec if avg_batch_time_sec is not None else "",
            "max_batch_time_sec": max_batch_time_sec if max_batch_time_sec is not None else "",
            "dataset_share_prep_time_sec": dataset_share_prep_time_sec if dataset_share_prep_time_sec is not None else "",
            "training_time_sec": training_time_sec if training_time_sec is not None else "",
            "client_distribution_time_sec": client_metrics.get("client_distribution_time_sec", ""),
            "client_eval_time_sec": client_metrics.get("client_eval_time_sec", ""),
            "client_eval_upload_time_sec": client_eval_upload_time_sec if client_eval_upload_time_sec is not None else "",
            "memory_baseline_mb": total_base_mb if total_base_mb is not None else "",
            "memory_peak_mb": total_peak_mb if total_peak_mb is not None else "",
            "memory_overhead_mb": total_overhead_mb if total_overhead_mb is not None else "",
            "node_memory_baseline_mb": json.dumps({k: round(v / mb, 3) for k, v in node_mem_baseline.items()}),
            "node_memory_peak_mb": json.dumps({k: round(v / mb, 3) for k, v in node_mem_peak.items()}),
            "node_memory_overhead_mb": json.dumps({k: round(v / mb, 3) for k, v in node_mem_overhead.items()}),
            "node_return_codes": json.dumps(return_codes),
>>>>>>> origin/main
            "log_dir": str(run_dir),
            "command": " ".join(sys.argv),
        }
        recorded = _append_run_to_xlsx(str(args.record_results_xlsx), row)
        print(f"  recorded_to_xlsx: {recorded}")


<<<<<<< HEAD
def main() -> None:
    args = parse_and_validate_args()
    if bool(getattr(args, "record_results_xlsx", "")) and not bool(getattr(args, "headless", False)):
=======
def main():
    args = parse_and_validate_args()
        
    print("=" * 70)
    print("Starting SENTRA Multi-Node Training")
    print("=" * 70)
    print(f"Total nodes: {args.n_nodes}")
    print(f"Base port: {args.base_port}")
    print(f"Host: {args.host}")
    print(f"Training: epochs={args.num_epochs}, batch_size={args.batch_size}, lr={args.learning_rate}")
    if args.batched:
        print(f"Batched secure config: field={args.field_size}, scale={args.scale_factor}, temp={args.softmax_temperature}, grad_clip={args.grad_clip}, logit_clip={args.logit_clip}, exp={args.exp_approx}, grad_mode={args.softmax_grad_mode}, loss_mode={args.loss_mode}")
        print(f"Batch config: batch_size={args.batch_size}, accum_steps={args.accum_steps}, effective_batch={args.batch_size * args.accum_steps}")
        if args.distribute_dataset_shares:
            print(f"Dataset sharing mode: owner-node (owner={args.dataset_owner_node})")
        elif args.receive_dataset_shares_from_client:
            print(f"Dataset sharing mode: external client sender (source_node_id={args.dataset_source_node_id})")
        if args.client_eval_after_training:
            print(f"Client-side eval after training: enabled ({args.client_eval_samples} samples)")
        print(f"Membership epoch: e={args.membership_epoch} (prefix m{args.membership_epoch}_)")
        print(f"Instability thresholds: logit={args.explode_logit_threshold}, loss_growth={args.loss_growth_threshold}, grad_norm={args.grad_norm_threshold}, abort={not args.no_abort_on_instability}")
        print(
            f"Debug flags: numerics={args.debug_numerics}, division={args.debug_division}, "
            f"packed_forward_pilot={args.packed_forward_pilot}, "
            f"packed_forward_native={args.packed_forward_native}, "
            f"packed_end2end={args.packed_end2end}"
        )
    print(f"Thresholds: t={args.t}")
    print(f"Dataset: {args.dataset}")
    print("=" * 70)
    print()

    if args.record_results_xlsx and not args.headless:
>>>>>>> origin/main
        print("Enabling --headless because --record-results-xlsx was provided.")
        args.headless = True
    if bool(getattr(args, "headless", False)):
        _run_headless(args)
        return
    raise SystemExit("Non-headless mode is deprecated; use --headless.")


if __name__ == "__main__":
    main()

