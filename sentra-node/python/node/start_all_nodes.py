"""SENTRA launcher (MNIST-only, batched secure runner).

Canonical multi-node entrypoint:
  start_all_nodes.py -> run_mnist_batched_secure.py (per node)

CLI parsing/validation lives in `start_all_nodes_cli.py`.
"""

from __future__ import annotations

import datetime
import os
import re
import subprocess
import sys
import time
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

_THIS_DIR = Path(__file__).resolve().parent
if str(_THIS_DIR) not in sys.path:
    sys.path.insert(0, str(_THIS_DIR))

from start_all_nodes_cli import parse_and_validate_args

ProcEntry = Tuple[int, subprocess.Popen]

_NODE_ROOT = Path(__file__).resolve().parent
_CLIENT_ROOT = _NODE_ROOT.parent / "client"


def _resolve_from_node_root(path_like: str) -> str:
    """Resolve relative paths against node root for cross-process consistency."""
    p = Path(str(path_like))
    if p.is_absolute():
        return str(p)
    return str((_NODE_ROOT / p).resolve())


def _quiet_tf_env(base: Optional[Dict[str, str]] = None) -> Dict[str, str]:
    """Default TF log level so headless runs do not fill stderr with repeated oneDNN banners."""
    e = dict(base if base is not None else os.environ)
    # 0=all, 1=hide INFO, 2=hide WARNING, 3=errors only. User can override in environment.
    e.setdefault("TF_CPP_MIN_LOG_LEVEL", "2")
    return e


def _logs_root() -> Path:
    return _NODE_ROOT / "logs"


def _safe_read_log(path: Path) -> str:
    if not path.is_file():
        return ""
    try:
        return path.read_text(encoding="utf-8", errors="ignore")
    except OSError as e:
        print(f"[headless] warning: could not read log {path}: {e}")
        return ""


def _first_timing(timings: Dict[str, float], *keys: str) -> Optional[float]:
    for key in keys:
        if key in timings:
            return float(timings[key])
    return None


def _format_timing_table(headers: List[str], rows: List[List[Any]]) -> str:
    def _cell(value: Any) -> str:
        if value is None:
            return "-"
        if isinstance(value, float):
            return f"{value:.2f}"
        return str(value)

    str_rows = [[_cell(c) for c in row] for row in rows]
    widths = [
        max(len(headers[i]), *(len(r[i]) for r in str_rows), 3)
        for i in range(len(headers))
    ]
    lines = [
        "  "
        + "  ".join(headers[i].ljust(widths[i]) for i in range(len(headers))),
        "  "
        + "  ".join("-" * widths[i] for i in range(len(headers))),
    ]
    for row in str_rows:
        lines.append(
            "  " + "  ".join(row[i].rjust(widths[i]) for i in range(len(headers)))
        )
    return "\n".join(lines)


def _print_run_timing_tables(
    timings: Dict[str, float], *, n_nodes: int, client_present: bool
) -> None:
    node_headers = [
        "node",
        "cold_start",
        "dataset",
        "training",
        "prover",
        "eval_upload",
        "total",
    ]
    node_rows: List[List[Any]] = []
    for nid in range(1, int(n_nodes) + 1):
        prefix = f"node{nid}_"
        node_rows.append(
            [
                nid,
                _first_timing(timings, f"{prefix}cold_start_sec"),
                _first_timing(
                    timings,
                    f"{prefix}dataset_prep_sec",
                    f"{prefix}dataset_sec",
                ),
                _first_timing(timings, f"{prefix}training_sec"),
                _first_timing(timings, f"{prefix}prover_sec"),
                _first_timing(timings, f"{prefix}client_eval_upload_sec"),
                _first_timing(timings, f"{prefix}total_sec"),
            ]
        )
    if any(any(v is not None for v in row[1:]) for row in node_rows):
        print("\n  Node timings (sec):")
        print(_format_timing_table(node_headers, node_rows))

    if client_present:
        client_headers = ["role", "cold_start", "distribution", "eval", "eval_wait", "eval_reconstruct", "total"]
        client_row = [
            "client",
            _first_timing(timings, "client_cold_start_sec"),
            _first_timing(
                timings,
                "client_distribution_sec",
                "client_client_distribution_sec",
            ),
            _first_timing(timings, "client_eval_sec", "client_client_eval_sec"),
            _first_timing(timings, "client_eval_share_wait_sec", "client_client_eval_share_wait_sec"),
            _first_timing(timings, "client_eval_reconstruct_sec", "client_client_eval_reconstruct_sec"),
            _first_timing(timings, "client_total_sec"),
        ]
        if any(v is not None for v in client_row[1:]):
            print("\n  Client timings (sec):")
            print(_format_timing_table(client_headers, [client_row]))


def _print_run_overhead_summary(overhead: Dict[str, float]) -> None:
    print("\n  Orchestrator timing (sec):")
    labels = [
        ("end_to_end_sec", "end_to_end"),
        ("pre_spawn_orchestration_sec", "pre_spawn_orchestration"),
        ("parallel_run_sec", "parallel_run"),
        ("spawn_to_last_node_exit_sec", "spawn_to_last_node_exit"),
        ("max_cold_start_sec", "max_cold_start"),
        ("max_mpc_work_sec", "max_mpc_work"),
        ("critical_path_logged_sec", "critical_path_logged"),
        ("post_log_process_exit_sec", "post_log_process_exit"),
        ("unaccounted_orchestrator_sec", "unaccounted_orchestrator"),
    ]
    for key, label in labels:
        if key in overhead:
            print(f"    {label}: {float(overhead[key]):.2f}")


def _print_log_tail(path: Path, *, label: str, max_lines: int = 40) -> None:
    text = _safe_read_log(path)
    if not text.strip():
        print(f"[headless] {label}: (no log at {path})")
        return
    lines = text.splitlines()
    tail = lines[-max_lines:] if len(lines) > max_lines else lines
    print(f"[headless] --- {label} (last {len(tail)} lines) ---")
    for line in tail:
        print(line)


def _env_with_node_on_pythonpath() -> Dict[str, str]:
    """Allow ``python -m sentra_client`` from ``client/`` to import ``ml_training`` from ``node/``."""
    e = _quiet_tf_env()
    node_path = str(_NODE_ROOT)
    prev = (e.get("PYTHONPATH") or "").strip()
    e["PYTHONPATH"] = node_path if not prev else f"{node_path}{os.pathsep}{prev}"
    return e


def _build_node_command(args, node_id: int) -> List[str]:
    cmd: List[str] = [
        sys.executable,
        "-u",
        str(_NODE_ROOT / "run_mnist_batched_secure.py"),
        "--node-id",
        str(int(node_id)),
        "--n-nodes",
        str(int(args.n_nodes)),
        "--t",
        str(int(args.t)),
        "--enable-network",
        "--base-port",
        str(int(args.base_port)),
        "--host",
        str(args.host),
        "--batch-size",
        str(int(args.batch_size)),
        "--accum-steps",
        str(int(args.accum_steps)),
        "--num-epochs",
        str(int(args.num_epochs)),
        "--learning-rate",
        str(float(args.learning_rate)),
        "--mnist-samples",
        str(int(args.mnist_samples)),
        "--seed",
        str(int(args.seed)),
        "--loss-mode",
        str(args.loss_mode),
        "--scale-factor",
        str(int(args.scale_factor)),
        "--field-size",
        str(int(args.field_size)),
        "--softmax-temperature",
        str(float(args.softmax_temperature)),
        "--grad-clip",
        str(float(args.grad_clip)),
        "--logit-clip",
        str(float(args.logit_clip)),
        "--exp-approx",
        str(args.exp_approx),
        "--softmax-grad-mode",
        str(args.softmax_grad_mode),
        "--explode-logit-threshold",
        str(float(args.explode_logit_threshold)),
        "--loss-growth-threshold",
        str(float(args.loss_growth_threshold)),
        "--grad-norm-threshold",
        str(float(args.grad_norm_threshold)),
        "--membership-epoch",
        str(int(args.membership_epoch)),
    ]

    if bool(args.no_abort_on_instability):
        cmd.append("--no-abort-on-instability")
    if bool(args.debug_numerics):
        cmd.append("--debug-numerics")
    if bool(args.debug_division):
        cmd.append("--debug-division")
    if bool(args.packed_forward_pilot):
        cmd.append("--packed-forward-pilot")
    if bool(args.packed_forward_native):
        cmd.append("--packed-forward-native")
    if bool(args.packed_end2end):
        cmd.append("--packed-end2end")
    if int(args.dpss_refresh_interval) > 0:
        cmd.extend(["--dpss-refresh-interval", str(int(args.dpss_refresh_interval))])

    if bool(args.enable_failure_detection):
        cmd.append("--enable-failure-detection")
    if bool(getattr(args, "enable_dropout_reshare_recovery", False)):
        cmd.append("--enable-dropout-reshare-recovery")
    if bool(getattr(args, "enable_join_recovery", False)):
        cmd.append("--enable-join-recovery")

    if bool(args.distribute_dataset_shares):
        cmd.extend(
            [
                "--distribute-dataset-shares",
                "--dataset-owner-node",
                str(int(args.dataset_owner_node)),
                "--dataset-distribution-timeout",
                str(float(args.dataset_distribution_timeout)),
            ]
        )
    if bool(args.receive_dataset_shares_from_client):
        cmd.extend(
            [
                "--receive-dataset-shares-from-client",
                "--dataset-source-node-id",
                str(int(args.dataset_source_node_id)),
                "--dataset-distribution-timeout",
                str(float(args.dataset_distribution_timeout)),
            ]
        )
        if bool(args.client_eval_after_training):
            cmd.extend(
                [
                    "--client-eval-after-training",
                    "--client-eval-samples",
                    str(int(args.client_eval_samples)),
                ]
            )
            if bool(getattr(args, "client_eval_batched_receive", False)):
                cmd.append("--client-eval-batched-receive")

    if bool(args.export_reconstructed_model):
        cmd.extend(
            [
                "--export-reconstructed-model",
                str(args.export_reconstructed_model),
                "--export-timeout",
                str(float(args.export_timeout)),
            ]
        )
    if bool(getattr(args, "init_weights_from_npz", "")):
        cmd.extend(
            [
                "--init-weights-from-npz",
                str(args.init_weights_from_npz),
            ]
        )
    if bool(getattr(args, "client_shares_model_weights", False)):
        cmd.extend(
            [
                "--receive-weights-shares-from-client",
                "--weights-source-node-id",
                str(int(getattr(args, "dataset_source_node_id", 0))),
            ]
        )

    if bool(getattr(args, "use_kvs_dataset", False)):
        cmd.append("--use-kvs-dataset")
    if bool(getattr(args, "use_weight_versioning", False)):
        cmd.append("--use-weight-versioning")

    return cmd


def _extract_metrics_from_log(log_text: str) -> Dict[str, Any]:
    metrics: Dict[str, Any] = {}
    acc = re.findall(r"Epoch\s+(\d+)\s+Test Accuracy:\s*([0-9eE+.\-]+)%", log_text)
    loss = re.findall(r"Epoch\s+(\d+)\s+Test Loss:\s*([0-9eE+.\-]+)", log_text)
    if acc:
        _ep, _val = acc[-1]
        metrics["final_epoch_acc_pct"] = float(_val)
    if loss:
        _ep, _val = loss[-1]
        metrics["final_epoch_loss"] = float(_val)
    return metrics


def _append_run_to_xlsx(path: str, row: Dict[str, Any]) -> str:
    try:
        from openpyxl import Workbook, load_workbook  # type: ignore
    except Exception:
        return path
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    if p.exists():
        wb = load_workbook(p)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.append(list(row.keys()))
    ws.append([row.get(k, "") for k in row.keys()])
    wb.save(p)
    return str(p)


def _run_headless(args) -> None:
    in_occlum = os.environ.get("SENTRA_IN_OCCLUM") == "1"
    if in_occlum and int(args.n_nodes) > 5:
        print(
            "[headless] NOTE: SGX/Occlum all-in-one runs many party processes inside one "
            "enclave (~5.4GB EPC). If startup fails, reduce n-nodes or use sentra-deployment "
            "multi-node SGX (one container per party)."
        )

    run_id = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    run_dir = _logs_root() / f"run_{run_id}"
    run_dir.mkdir(parents=True, exist_ok=True)

    t_run_start = time.time()
    procs: List[ProcEntry] = []
    open_files = []
    for nid in range(1, int(args.n_nodes) + 1):
        out_path = run_dir / f"node_{nid}.log"
        out_f = open(out_path, "w", encoding="utf-8")
        open_files.append(out_f)
        p = subprocess.Popen(
            _build_node_command(args, nid),
            cwd=str(_NODE_ROOT),
            stdout=out_f,
            stderr=subprocess.STDOUT,
            env=_quiet_tf_env(),
        )
        procs.append((nid, p))
        print(f"[headless] started node {nid} -> {out_path}")
        time.sleep(1.0 if in_occlum else 0.4)

    client_proc: Optional[subprocess.Popen] = None
    client_log = run_dir / "client_distributor.log"
    if bool(getattr(args, "start_client_distributor", False)):
        if in_occlum:
            print("[headless] waiting for nodes to bind before starting client...")
            time.sleep(8.0)
        cf = open(client_log, "w", encoding="utf-8")
        open_files.append(cf)
        cmd = [
            sys.executable,
            "-u",
            "-m",
            "sentra_client",
            "--client-node-id",
            str(int(getattr(args, "dataset_source_node_id", 0))),
            "--n-nodes",
            str(int(args.n_nodes)),
            "--t",
            str(int(args.t)),
            "--base-port",
            str(int(args.base_port)),
            "--host",
            str(args.host),
            "--mnist-samples",
            str(int(args.mnist_samples)),
            "--client-test-samples",
            str(int(args.client_test_samples)),
            "--field-size",
            str(int(args.field_size)),
            "--scale-factor",
            str(int(args.scale_factor)),
            "--seed",
            str(int(args.seed)),
            "--barrier-timeout",
            str(float(args.dataset_distribution_timeout)),
            "--membership-epoch",
            str(int(args.membership_epoch)),
        ]
        if bool(getattr(args, "infer_image", "")):
            cmd.extend(["--infer-image", _resolve_from_node_root(str(args.infer_image))])
            if bool(getattr(args, "infer_raw_mnist", False)):
                cmd.append("--infer-raw-mnist")
        if bool(getattr(args, "init_weights_from_npz", "")):
            cmd.extend(["--init-weights-from-npz", _resolve_from_node_root(str(args.init_weights_from_npz))])
        if bool(getattr(args, "client_shares_model_weights", False)):
            cmd.append("--share-model-weights")
        if bool(args.client_eval_after_training):
            cmd.extend(
                [
                    "--collect-client-eval",
                    "--client-eval-samples",
                    str(int(args.client_eval_samples)),
                    "--eval-timeout",
                    str(float(args.client_eval_timeout)),
                ]
            )
            if bool(getattr(args, "client_eval_batched_receive", False)):
                cmd.append("--client-eval-batched-receive")
        client_env = _env_with_node_on_pythonpath()
        if in_occlum:
            client_env["SENTRA_IN_OCCLUM"] = "1"
        client_proc = subprocess.Popen(
            cmd, cwd=str(_CLIENT_ROOT), env=client_env, stdout=cf, stderr=cf
        )
        print(f"[headless] started client distributor -> {client_log}")

    t_parallel_start = time.time()
    pre_spawn_orchestration_sec = float(t_parallel_start - t_run_start)
    return_codes: Dict[int, int] = {}
    t_last_node_exit: Optional[float] = None
    for nid, p in procs:
        p.wait()
        return_codes[nid] = int(p.returncode or 0)
        t_last_node_exit = time.time()
        print(f"[headless] node {nid} exited with code {return_codes[nid]}")
        if return_codes[nid] != 0:
            _print_log_tail(run_dir / f"node_{nid}.log", label=f"node {nid} log")
    spawn_to_last_node_exit_sec = float(
        (t_last_node_exit or t_parallel_start) - t_parallel_start
    )
    if client_proc is not None:
        client_timeout = float(getattr(args, "client_eval_timeout", 0.0) or 0.0)
        if client_timeout <= 0:
            client_timeout = float(getattr(args, "dataset_distribution_timeout", 900.0))
        try:
            client_proc.wait(timeout=max(client_timeout + 60.0, 900.0))
        except subprocess.TimeoutExpired:
            print("[headless] warning: client distributor did not exit before timeout")
            client_proc.kill()
        except Exception:
            pass
    t_run_end = time.time()
    end_to_end_sec = float(t_run_end - t_run_start)
    parallel_run_sec = float(t_run_end - t_parallel_start)

    for f in open_files:
        try:
            f.close()
        except Exception:
            pass

    status = "success" if all(c == 0 for c in return_codes.values()) else "failed"

    node1_text = _safe_read_log(run_dir / "node_1.log")
    m1 = _extract_metrics_from_log(node1_text)
    client_text = _safe_read_log(client_log)
    if status != "success" and not node1_text.strip():
        print(
            f"[headless] Log files may be on the Docker volume at "
            f"{(_logs_root()).resolve()} (host mount: /workspace/node/logs)."
        )
    client_acc = None
    m_client = re.findall(r"Client Final Accuracy \(\d+ samples\):\s*([0-9eE+.\-]+)%", client_text)
    if m_client:
        client_acc = float(m_client[-1])

    client_eval = bool(getattr(args, "client_eval_after_training", False)) and client_proc is not None

    node_logs: Dict[int, str] = {}
    for nid in return_codes:
        node_logs[int(nid)] = _safe_read_log(run_dir / f"node_{nid}.log")

    from ml_training.benchmark_stats import build_run_overhead_metrics, summarize_run_logs

    bench = summarize_run_logs(
        node_logs=node_logs,
        client_log=client_text,
        wall_sec=end_to_end_sec,
    )
    timings: Dict[str, float] = bench.get("timings") or {}
    overhead = build_run_overhead_metrics(
        timings,
        end_to_end_sec=end_to_end_sec,
        pre_spawn_orchestration_sec=pre_spawn_orchestration_sec,
        spawn_to_last_node_exit_sec=spawn_to_last_node_exit_sec,
        parallel_run_sec=parallel_run_sec,
        n_nodes=int(args.n_nodes),
        client_present=client_proc is not None,
    )

    print("\nRun summary:")
    print(f"  status: {status}")
    if client_eval:
        if client_acc is not None:
            print(f"  final_accuracy_pct: {client_acc:.2f}")
        else:
            print("  final_accuracy_pct: (not found in client_distributor.log)")
    else:
        print(f"  final_accuracy_pct: {m1.get('final_epoch_acc_pct', 0.0)}")
        if m1.get("final_epoch_loss") is not None:
            print(f"  final_epoch_loss: {m1.get('final_epoch_loss')}")
    print(f"  end_to_end_sec: {end_to_end_sec:.2f}")
    print(f"  wall_clock_sec: {spawn_to_last_node_exit_sec:.2f}  (spawn -> last node exit)")
    _print_run_overhead_summary(overhead)
    _print_run_timing_tables(
        timings,
        n_nodes=int(args.n_nodes),
        client_present=client_proc is not None,
    )
    print(f"  logs: {run_dir}")

    if bool(getattr(args, "record_results_xlsx", "")):
        row = {
            "timestamp": datetime.datetime.now().isoformat(timespec="seconds"),
            "status": status,
            "duration_sec": end_to_end_sec,
            "spawn_to_last_node_exit_sec": spawn_to_last_node_exit_sec,
            "max_cold_start_sec": overhead.get("max_cold_start_sec", ""),
            "critical_path_logged_sec": overhead.get("critical_path_logged_sec", ""),
            "n_nodes": int(args.n_nodes),
            "dataset": str(args.dataset),
            "train_mode": "secure",
            "loss_mode": str(args.loss_mode),
            "softmax_grad_mode": str(args.softmax_grad_mode),
            "num_epochs": int(args.num_epochs),
            "batch_size": int(args.batch_size),
            "learning_rate": float(args.learning_rate),
            "field_size": int(args.field_size),
            "scale_factor": int(args.scale_factor),
            "softmax_temperature": float(args.softmax_temperature),
            "seed": int(args.seed),
            "final_epoch_acc_pct": (
                client_acc if client_eval and client_acc is not None else m1.get("final_epoch_acc_pct", "")
            ),
            "final_epoch_loss": m1.get("final_epoch_loss", ""),
            "reconstructed_acc": client_acc if client_acc is not None else "",
            "log_dir": str(run_dir),
            "command": " ".join(sys.argv),
        }
        recorded = _append_run_to_xlsx(str(args.record_results_xlsx), row)
        print(f"  recorded_to_xlsx: {recorded}")


def main() -> None:
    args = parse_and_validate_args()
    if bool(getattr(args, "record_results_xlsx", "")) and not bool(getattr(args, "headless", False)):
        print("Enabling --headless because --record-results-xlsx was provided.")
        args.headless = True
    if bool(getattr(args, "headless", False)):
        _run_headless(args)
        return
    raise SystemExit("Non-headless mode is deprecated; use --headless.")


if __name__ == "__main__":
    main()
