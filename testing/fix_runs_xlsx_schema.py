"""
Repair/migrate SENTRA run workbook headers to the current schema.

Usage:
  python testing/fix_runs_xlsx_schema.py --in logs/sentra_runs.xlsx --out logs/sentra_runs_fixed.xlsx
"""

import argparse
from pathlib import Path
from openpyxl import load_workbook, Workbook


NEW_HEADERS = [
    "timestamp",
    "run_start",
    "run_end",
    "run_id",
    "status",
    "duration_sec",
    "duration_min",
    "n_nodes",
    "dataset",
    "batched",
    "train_mode",
    "loss_mode",
    "softmax_grad_mode",
    "num_epochs",
    "batch_size",
    "learning_rate",
    "field_size",
    "scale_factor",
    "softmax_temperature",
    "grad_clip",
    "logit_clip",
    "seed",
    "final_epoch_acc_pct",
    "final_epoch_loss",
    "reconstructed_acc",
    "node_success_count",
    "node_elapsed_sec",
    "prover_time_sec",
    "avg_batch_time_sec",
    "max_batch_time_sec",
    "memory_baseline_mb",
    "memory_peak_mb",
    "memory_overhead_mb",
    "node_memory_baseline_mb",
    "node_memory_peak_mb",
    "node_memory_overhead_mb",
    "node_return_codes",
    "log_dir",
    "command",
]

OLD_HEADERS = [
    "timestamp",
    "run_id",
    "status",
    "duration_sec",
    "n_nodes",
    "dataset",
    "batched",
    "train_mode",
    "loss_mode",
    "softmax_grad_mode",
    "num_epochs",
    "batch_size",
    "learning_rate",
    "field_size",
    "scale_factor",
    "softmax_temperature",
    "grad_clip",
    "logit_clip",
    "seed",
    "final_epoch_acc_pct",
    "final_epoch_loss",
    "reconstructed_acc",
    "node_success_count",
    "node_return_codes",
    "log_dir",
    "command",
]


def _is_run_id(v):
    return isinstance(v, str) and v.startswith("run_")


def _row_to_new_dict(row_vals):
    d = {}
    # Old schema row: run_id in col2
    if len(row_vals) >= 2 and _is_run_id(row_vals[1]):
        for i, h in enumerate(OLD_HEADERS):
            d[h] = row_vals[i] if i < len(row_vals) else None
        return d
    # New schema row (written despite old header): run_id in col4
    if len(row_vals) >= 4 and _is_run_id(row_vals[3]):
        for i, h in enumerate(NEW_HEADERS):
            d[h] = row_vals[i] if i < len(row_vals) else None
        return d
    # Fallback: map by new positional layout
    for i, h in enumerate(NEW_HEADERS):
        d[h] = row_vals[i] if i < len(row_vals) else None
    return d


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--in", dest="in_path", required=True, help="Input workbook path")
    ap.add_argument("--out", dest="out_path", required=True, help="Output workbook path")
    args = ap.parse_args()

    in_path = Path(args.in_path)
    out_path = Path(args.out_path)

    wb_in = load_workbook(in_path)
    ws_in = wb_in.active

    wb_out = Workbook()
    ws_out = wb_out.active
    ws_out.title = "runs"
    ws_out.append(NEW_HEADERS)

    max_col = ws_in.max_column
    migrated = 0
    for r in range(2, ws_in.max_row + 1):
        row_vals = [ws_in.cell(r, c).value for c in range(1, max_col + 1)]
        if all(v is None for v in row_vals):
            continue
        d = _row_to_new_dict(row_vals)
        ws_out.append([d.get(h, "") for h in NEW_HEADERS])
        migrated += 1

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb_out.save(out_path)
    print(f"Migrated rows: {migrated}")
    print(f"Written: {out_path}")


if __name__ == "__main__":
    main()

